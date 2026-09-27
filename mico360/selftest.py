"""Post-build self-test: ``MICO360DocToolkit.exe --selftest <report.json>``.

A frozen (PyInstaller) app can't run the Python test suite, so this mode
exercises the *packaged* app from the inside: every bundled library is
imported, every major tool is run on generated sample files, OCR runs on a
rendered scan, and the whole UI is built offscreen. It proves the installer
contains everything the app needs.

It never touches the user's data: settings go to a throw-away INI, samples and
outputs live in a temp folder that is deleted afterwards, nothing is
downloaded (the on-demand LibreOffice engine and Microsoft Word are bypassed,
so Office → PDF uses the built-in renderer), and no registry keys are written.
The result is written as JSON (plus a readable .txt next to it); the exit code
is 0 only when every check passed.
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
import time
import traceback
from pathlib import Path

_IMPORTS = [
    "fitz", "pypdf", "PIL", "pillow_heif", "pptx", "docx", "openpyxl", "pdfplumber",
    "pdfminer", "pdf2docx", "reportlab", "cryptography", "cv2", "numpy",
    "onnxruntime", "rapidocr_onnxruntime", "vtracer", "shapely", "pyclipper",
    "PySide6.QtNetwork", "PySide6.QtSvg",
]
_WIN_IMPORTS = ["win32file", "win32con", "win32security", "pythoncom", "pywintypes"]


class _Report:
    """The ``report`` callable processors expect (collects log lines)."""

    def __init__(self):
        self.lines: list[str] = []

    def __call__(self, msg, *a, **k):
        self.lines.append(str(msg))


def _samples(d: Path) -> dict:
    import fitz
    from PIL import Image, ImageDraw, ImageFont
    s: dict = {}
    # A 3-page PDF with bookmarks and metadata.
    pdf = d / "sample.pdf"
    doc = fitz.open()
    for i in range(3):
        pg = doc.new_page()
        pg.insert_text((72, 90), f"MICO360 self-test page {i + 1}", fontsize=18)
        pg.insert_textbox(fitz.Rect(72, 120, 520, 700), "Lorem ipsum dolor sit amet. " * 30,
                          fontsize=10)
    doc.set_toc([[1, "Start", 1], [1, "End", 3]])
    doc.set_metadata({"title": "Self-test", "author": "MICO360"})
    doc.save(str(pdf)); doc.close()
    s["pdf"] = pdf
    pdf2 = d / "second.pdf"
    doc = fitz.open(); doc.new_page().insert_text((72, 90), "Second file"); doc.save(str(pdf2))
    doc.close()
    s["pdf2"] = pdf2
    # A table PDF for PDF → Excel.
    tbl = d / "table.pdf"
    doc = fitz.open(); pg = doc.new_page()
    for r in range(4):
        for c in range(3):
            rect = fitz.Rect(72 + c * 120, 100 + r * 30, 192 + c * 120, 130 + r * 30)
            pg.draw_rect(rect)
            pg.insert_textbox(rect + (4, 6, -4, -4), f"R{r}C{c}", fontsize=10)
    doc.save(str(tbl)); doc.close()
    s["table_pdf"] = tbl
    # Images: a photo-like JPEG, a transparent PNG, a HEIC.
    img = Image.new("RGB", (800, 600), (40, 90, 160))
    dr = ImageDraw.Draw(img)
    dr.rectangle((100, 100, 700, 500), fill=(230, 200, 60))
    jpg = d / "photo.jpg"; img.save(jpg, quality=92)
    png = d / "logo.png"
    Image.new("RGBA", (300, 200), (200, 30, 30, 180)).save(png)
    s.update(jpg=jpg, png=png)
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
        heic = d / "phone.heic"
        img.save(heic, "HEIF", quality=90)
        s["heic"] = heic
    except Exception:
        pass
    # A "scanned" PDF (image only) with large clear text for OCR.
    scan_img = Image.new("RGB", (1700, 600), "white")
    dr = ImageDraw.Draw(scan_img)
    try:
        font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 96)
    except Exception:
        font = ImageFont.load_default()
    dr.text((80, 220), "INVOICE TOTAL 1250", fill="black", font=font)
    scan_png = d / "scan.png"; scan_img.save(scan_png)
    scan = d / "scan.pdf"
    doc = fitz.open(); pg = doc.new_page(width=850, height=300)
    pg.insert_image(pg.rect, filename=str(scan_png)); doc.save(str(scan)); doc.close()
    s["scan_pdf"] = scan
    # SVG, Word, Excel.
    svg = d / "shape.svg"
    svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="120">'
                   '<rect x="10" y="10" width="180" height="100" fill="#8a1c2b"/></svg>',
                   encoding="utf-8")
    s["svg"] = svg
    import docx
    wd = docx.Document()
    wd.add_heading("Self-test document", 1)
    wd.add_paragraph("A paragraph with **text** for conversion.")
    t = wd.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "A"; t.cell(1, 1).text = "D"
    word = d / "letter.docx"; wd.save(word)
    s["docx"] = word
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["Item", "Qty"]); ws.append(["Pens", 3])
    xlsx = d / "sheet.xlsx"; wb.save(xlsx)
    s["xlsx"] = xlsx
    return s


def _run_tool(tid: str, src, out_dir: Path, overrides: dict | None = None) -> list:
    from mico360.core.tools import TOOLS_BY_ID
    tool = TOOLS_BY_ID[tid]
    opt = {o.key: o.default for o in tool.options}
    opt.update(overrides or {})
    opt.setdefault("overwrite", True)
    rep = _Report()
    outs = tool.runner(src, out_dir, opt, rep) or []
    outs = [Path(o) for o in outs]
    empty = [o.name for o in outs if not (o.exists() and o.stat().st_size > 0)]
    if not outs:
        raise AssertionError("no output produced")
    if empty:
        raise AssertionError(f"empty output: {empty}")
    return outs


def run(report_path: str) -> int:
    results: list[dict] = []
    t_all = time.monotonic()
    work = Path(tempfile.mkdtemp(prefix="mico360_selftest_"))
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

    def check(name, fn):
        t0 = time.monotonic()
        try:
            detail = fn()
            results.append({"check": name, "ok": True, "seconds": round(time.monotonic() - t0, 2),
                            "detail": str(detail or "")[:300]})
        except Exception as exc:  # noqa: BLE001 - every failure is reported
            results.append({"check": name, "ok": False, "seconds": round(time.monotonic() - t0, 2),
                            "detail": f"{type(exc).__name__}: {exc}"[:500],
                            "trace": traceback.format_exc()[-1500:]})

    # Settings: a throw-away INI, and nothing is ever downloaded.
    from PySide6.QtCore import QSettings
    from mico360.config import settings
    settings._s = QSettings(str(work / "DocToolkit.ini"), QSettings.IniFormat)
    settings.auto_download_engine = False
    settings.ai_enabled = False

    import importlib
    for mod in _IMPORTS + (_WIN_IMPORTS if sys.platform.startswith("win") else []):
        check(f"import {mod}", lambda m=mod: importlib.import_module(m) and "ok")

    from mico360 import __version__
    check("app version", lambda: __version__)

    samples: dict = {}

    def _make_samples():
        (work / "in").mkdir(parents=True, exist_ok=True)
        samples.update(_samples(work / "in"))
        return ", ".join(sorted(samples))
    check("create sample files", _make_samples)

    from mico360.core import processors as P
    # Office → PDF: use the built-in renderer (no engine download, no Word automation).
    P.ensure_libreoffice = lambda *a, **k: None
    P._word_to_pdf_msword = lambda *a, **k: False

    out = work / "out"
    S = samples.get
    cases = [
        ("Compress PDF (lossless)", "pdf_compress", S("pdf"), {"level": "lossless"}),
        ("Compress PDF (medium)", "pdf_compress", S("pdf"), {"level": "medium"}),
        ("Compress PDF (target size)", "pdf_compress", S("pdf"), {"level": "target", "target_kb": 40}),
        ("Merge PDF", "pdf_merge", [S("pdf"), S("pdf2")], {"output_name": "merged"}),
        ("Split PDF", "pdf_split", S("pdf"), {"mode": "ranges", "ranges": "1-2, 3"}),
        ("Organize PDF (rotate)", "pdf_organize", S("pdf"), {"operation": "rotate", "angle": 90, "pages": "1"}),
        ("Protect PDF", "pdf_protect", S("pdf"), {"operation": "protect", "password": "pw123",
                                                  "confirm_password": "pw123"}),
        ("Watermark PDF", "pdf_watermark", S("pdf"), {"wm_type": "text", "text": "CONFIDENTIAL"}),
        ("Add Page Numbers", "pdf_page_numbers", S("pdf"), {}),
        ("Sign PDF", "pdf_sign", S("pdf"), {"image_path": str(S("png"))}),
        ("Edit Metadata", "pdf_metadata", S("pdf"), {"title": "Checked", "copyright": "(c) MICO360"}),
        ("PDF → Word", "pdf_convert", S("pdf"), {"target": "word"}),
        ("PDF → PowerPoint", "pdf_convert", S("pdf"), {"target": "pptx"}),
        ("PDF → Excel", "pdf_convert", S("table_pdf"), {"target": "excel"}),
        ("PDF → Images", "pdf_convert", S("pdf"), {"target": "image", "format": "png"}),
        ("Office → PDF (Word, built-in renderer)", "office_to_pdf", S("docx"), {}),
        ("Word → Markdown", "to_markdown", S("docx"), {}),
        ("Excel → Markdown", "to_markdown", S("xlsx"), {}),
        ("PDF → Markdown", "to_markdown", S("pdf"), {}),
        ("Image → PDF (combined)", "image_to_pdf", [S("jpg"), S("png")], {"combine": True,
                                                                         "output_name": "images"}),
        ("Compress Image", "image_compress", S("jpg"), {"level": "medium"}),
        ("Resize Image", "image_resize", S("jpg"), {"mode": "percent", "percent": 50}),
        ("Convert Image (PNG → WEBP)", "image_convert", S("png"), {"format": "webp"}),
        ("Watermark Image", "image_watermark", S("jpg"), {"wm_type": "text", "text": "MICO360"}),
        ("SVG → Image", "svg_to_image", S("svg"), {"format": "png"}),
        ("Image → SVG (trace)", "image_to_svg", S("png"), {"mode": "trace"}),
        ("Edit File Properties", "file_properties", S("pdf"), {"date_modified": "2026-01-02 10:30"}),
    ]
    if S("heic"):
        cases.append(("Convert Image (HEIC → JPG)", "image_convert", S("heic"), {"format": "jpg"}))
    for label, tid, src, ov in cases:
        if src is None or (isinstance(src, list) and None in src):
            results.append({"check": label, "ok": False, "seconds": 0,
                            "detail": "sample file missing"})
            continue
        check(label, lambda t=tid, s=src, o=ov, n=label: ", ".join(
            p.name for p in _run_tool(t, s, out / n.split(" ")[0].lower(), o)))

    def _unlock():
        prot = next((out / "protect").glob("*_protected.pdf"))
        return ", ".join(p.name for p in _run_tool(
            "pdf_protect", prot, out / "unlock", {"operation": "unlock", "password": "pw123"}))
    check("Unlock PDF", _unlock)

    def _ocr():
        import fitz
        o = _run_tool("pdf_ocr", S("scan_pdf"), out / "ocr", {"ocr_lang": "latin",
                                                               "ocr_use_ai": False})[0]
        text = fitz.open(str(o))[0].get_text()
        from mico360.core.processors import ocr_active_provider
        if "INVOICE" not in text.upper():
            raise AssertionError(f"OCR text layer missing the word: {text[:80]!r}")
        return f"text layer: {text.strip()[:40]!r} via {ocr_active_provider()}"
    check("Searchable PDF (OCR)", _ocr)

    def _ocr_repeat():
        """OCR keeps working after the first run: several different scans in a
        row (including large text), then re-running on the result."""
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        from mico360.core.util import ProcessError
        try:
            big = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 64)
        except Exception:
            big = ImageFont.load_default()
        words = ["RECEIPT page 1", "CONTRACT page 2", "INVOICE page 3"]
        outs = []
        for n, w in enumerate(words):
            im = Image.new("RGB", (1650, 2200), "white")
            ImageDraw.Draw(im).text((120, 200), w, fill="black", font=big)
            png = work / "in" / f"rep{n}.png"; im.save(png)
            pdf = work / "in" / f"rep{n}.pdf"
            d = fitz.open(); pg = d.new_page(width=612, height=816)
            pg.insert_image(pg.rect, filename=str(png)); d.save(str(pdf)); d.close()
            o = _run_tool("pdf_ocr", pdf, out / "ocr_repeat", {"ocr_use_ai": False,
                                                                "quality": "balanced"})[0]
            t = " ".join(fitz.open(str(o))[0].get_text().split())
            if w not in t:
                raise AssertionError(f"run {n + 1}: expected {w!r}, text layer {t!r}")
            outs.append(o)
        try:
            _run_tool("pdf_ocr", outs[0], out / "ocr_again", {"ocr_use_ai": False})
            raise AssertionError("re-OCR of a searchable PDF should explain there's nothing to do")
        except ProcessError:
            pass
        r = _run_tool("pdf_ocr", outs[0], out / "ocr_redo", {"ocr_use_ai": False,
                                                              "ocr_redo": True})[0]
        t = " ".join(fitz.open(str(r))[0].get_text().split())
        if t.count("RECEIPT") != 1:
            raise AssertionError(f"Redo OCR should replace, not double, the text: {t!r}")
        return f"{len(words)} scans in a row + re-run + redo OK"
    check("Searchable PDF (OCR) — repeated runs, large text, redo", _ocr_repeat)

    def _ocr_image():
        """An image (PNG) and a multi-page TIFF become searchable PDFs."""
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        try:
            font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 48)
        except Exception:
            font = ImageFont.load_default()

        def pg(text):
            im = Image.new("RGB", (1275, 1650), "white")
            ImageDraw.Draw(im).text((100, 160), text, fill="black", font=font)
            return im
        png = work / "in" / "photo_scan.png"
        pg("Delivery note 7731").save(png, dpi=(150, 150))
        o = _run_tool("pdf_ocr", png, out / "ocr_image", {"ocr_use_ai": False})[0]
        t = " ".join(fitz.open(str(o))[0].get_text().split())
        if o.suffix != ".pdf" or "Delivery note 7731" not in t:
            raise AssertionError(f"image OCR: {o.name} {t!r}")
        tif = work / "in" / "pages.tif"
        fr = [pg(f"Page number {n}") for n in (1, 2)]
        fr[0].save(tif, save_all=True, append_images=fr[1:], dpi=(150, 150))
        o2 = _run_tool("pdf_ocr", tif, out / "ocr_image", {"ocr_use_ai": False})[0]
        d = fitz.open(str(o2))
        texts = [" ".join(p.get_text().split()) for p in d]
        d.close()
        if len(texts) != 2 or "Page number 2" not in texts[1]:
            raise AssertionError(f"multi-page TIFF OCR: {texts}")
        return f"{o.name} + {o2.name} (2 pages)"
    check("Searchable PDF (OCR) from images (PNG, multi-page TIFF)", _ocr_image)

    def _ui():
        from PySide6.QtWidgets import QApplication
        app = QApplication.instance() or QApplication([])
        from mico360.core.tools import TOOLS
        from mico360.ui.main_window import MainWindow
        w = MainWindow()
        opened = 0
        for t in TOOLS:
            w.open_tool(t.id)
            page = w._current_tool_page()
            if page is None or page.tool.id != t.id:
                raise AssertionError(f"tool page didn't open: {t.id}")
            opened += 1
        w.open_settings()
        w.open_help()
        app.processEvents()
        w.close()
        return f"main window + {opened} tool pages + Settings + Help built"
    check("User interface (every page)", _ui)

    ok = all(r["ok"] for r in results)
    summary = {
        "app": "MICO360 Doc Toolkit", "version": __version__,
        "frozen": bool(getattr(sys, "frozen", False)), "platform": sys.platform,
        "passed": sum(r["ok"] for r in results), "failed": sum(not r["ok"] for r in results),
        "seconds": round(time.monotonic() - t_all, 1), "ok": ok, "results": results,
    }
    rp = Path(report_path)
    rp.parent.mkdir(parents=True, exist_ok=True)
    rp.write_text(json.dumps(summary, indent=2), encoding="utf-8")
    lines = [f"MICO360 Doc Toolkit v{__version__} self-test — "
             f"{summary['passed']} passed, {summary['failed']} failed "
             f"({summary['seconds']} s, frozen={summary['frozen']})", ""]
    for r in results:
        lines.append(f"[{'PASS' if r['ok'] else 'FAIL'}] {r['check']}  ({r['seconds']} s)"
                     + (f"  — {r['detail']}" if r["detail"] else ""))
    rp.with_suffix(".txt").write_text("\n".join(lines) + "\n", encoding="utf-8")
    try:
        import shutil
        shutil.rmtree(work, ignore_errors=True)
    except Exception:
        pass
    return 0 if ok else 1
